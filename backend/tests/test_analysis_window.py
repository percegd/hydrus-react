"""Testes da janela de análise e dias de leitura (T9).

Paridade com hydrus_core.input_validation.validate_analysis_window e com os
defaults do Gradio (início 1, fim 365, leitura "55, 128, 237, 365").
"""

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.schemas import SimulateRequest
from app.service import execute_job


def _series_payload(**overrides) -> dict:
    rain_series = [5.0 if day % 3 == 0 else 0.0 for day in range(365)]
    payload = {
        "soil_key": "latossolo_medio",
        "crop_key": "soja",
        "rain_series": rain_series,
        "solutes": [{"symbol": "K", "name": "Potássio", "dose_kg_ha": 60.0, "kd": 1.0}],
    }
    payload.update(overrides)
    return payload


def test_schema_defaults_match_gradio() -> None:
    req = SimulateRequest(soil_key="latossolo_medio")
    assert req.analysis_start_day == 1
    assert req.analysis_end_day == 365
    assert req.print_days is None


def test_window_start_after_end_returns_422() -> None:
    with TestClient(app) as client:
        response = client.post(
            "/simulate", json=_series_payload(analysis_start_day=200, analysis_end_day=100)
        )
    assert response.status_code == 422, response.text
    assert any(item.get("field") == "analysis_end_day" for item in response.json()["detail"])


def test_window_outside_calendar_returns_422() -> None:
    with TestClient(app) as client:
        response = client.post(
            "/simulate", json=_series_payload(analysis_start_day=0, analysis_end_day=400)
        )
    assert response.status_code == 422, response.text
    fields = {item.get("field") for item in response.json()["detail"]}
    assert "analysis_start_day" in fields
    assert "analysis_end_day" in fields


def test_print_day_outside_window_returns_422() -> None:
    with TestClient(app) as client:
        response = client.post(
            "/simulate",
            json=_series_payload(
                analysis_start_day=100, analysis_end_day=200, print_days=[50.0]
            ),
        )
    assert response.status_code == 422, response.text
    assert any(item.get("field") == "print_days_text" for item in response.json()["detail"])


def test_application_day_outside_window_returns_422() -> None:
    with TestClient(app) as client:
        response = client.post(
            "/simulate",
            json=_series_payload(
                analysis_start_day=100,
                analysis_end_day=200,
                placement={"application_day": 50.0},
            ),
        )
    assert response.status_code == 422, response.text
    assert any(item.get("field") == "app_day" for item in response.json()["detail"])


def test_empty_print_days_returns_422() -> None:
    with TestClient(app) as client:
        response = client.post("/simulate", json=_series_payload(print_days=[]))
    assert response.status_code == 422, response.text
    assert any(item.get("field") == "print_days_text" for item in response.json()["detail"])


def test_lamina_print_day_beyond_duration_returns_422() -> None:
    with TestClient(app) as client:
        response = client.post(
            "/simulate",
            json={
                "soil_key": "latossolo_medio",
                "total_days": 7,
                "rain_mm": 60.0,
                "solutes": [
                    {"symbol": "K", "name": "Potássio", "dose_kg_ha": 30.0, "kd": 0.45}
                ],
                "print_days": [55.0, 128.0],
            },
        )
    assert response.status_code == 422, response.text
    assert any(item.get("field") == "print_days_text" for item in response.json()["detail"])


def test_window_slices_series_and_report_labels() -> None:
    rain_series = [5.0 if day % 3 == 0 else 0.0 for day in range(365)]
    expected = sum(rain_series[99:200])
    with TestClient(app) as client:
        response = client.post(
            "/simulate",
            json=_series_payload(
                analysis_start_day=100,
                analysis_end_day=200,
                print_days=[100.0, 150.0, 200.0],
            ),
        )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["water_balance"]["applied_mm"] == expected
    # Janela parcial: rótulos de período (Gradio: Chuva/Drenagem no Período).
    assert body["water_balance"]["applied_label"] == "Chuva no Período"
    assert body["water_balance"]["drainage_label"] == "Drenagem no Período"


def test_window_and_print_days_reach_excel_and_report(tmp_path) -> None:
    rain_series = [5.0 if day % 3 == 0 else 0.0 for day in range(365)]
    req = SimulateRequest(
        soil_key="latossolo_medio",
        crop_key="soja",
        rain_series=rain_series,
        solutes=[],
        water_only=True,
        analysis_start_day=100,
        analysis_end_day=200,
        print_days=[100.0, 150.0, 200.0],
    )
    _resp, excel_path, report = execute_job(req, str(tmp_path))
    # Laudo sazonal com a janela efetiva (Gradio: "janela 100–200").
    assert "janela 100–200" in report
    # Dias de calendário convertidos p/ dias de ensaio (d − sw + 1): 1, 51, 101.
    workbook = load_workbook(excel_path)
    headers = [cell.value for cell in workbook["Perfil_Solo_Sazonal"][1]]
    umidade = [h for h in headers if str(h).startswith("Umidade_Dia_")]
    assert len(umidade) == 3
    assert "Umidade_Dia_1" in str(umidade[0])
    assert "Umidade_Dia_51" in str(umidade[1])
    assert "Umidade_Dia_101" in str(umidade[2])


def test_default_365_series_keeps_gradio_marks(tmp_path) -> None:
    rain_series = [5.0 if day % 3 == 0 else 0.0 for day in range(365)]
    req = SimulateRequest(
        soil_key="latossolo_medio",
        crop_key="soja",
        rain_series=rain_series,
        solutes=[],
        water_only=True,
    )
    _resp, excel_path, report = execute_job(req, str(tmp_path))
    assert "janela ano cheio (1–365)" in report
    workbook = load_workbook(excel_path)
    headers = [cell.value for cell in workbook["Perfil_Solo_Sazonal"][1]]
    umidade = [str(h) for h in headers if str(h).startswith("Umidade_Dia_")]
    for day in ("55", "128", "237", "365"):
        assert any(day in h for h in umidade), umidade
