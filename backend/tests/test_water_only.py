"""Testes do modo "Somente água" (T5) — paridade com o Gradio.

Espelha /workspace/hydrus/tests/test_water_only_mode.py: o laudo e os
artefatos derivados descrevem apenas o balanço hídrico, sem solutos.
"""

import base64
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.schemas import SimulateRequest
from app.service import execute_job

PNG_MAGIC = b"\x89PNG"


def _water_only_request(**overrides) -> SimulateRequest:
    data = {
        "soil_key": "latossolo_medio",
        "crop_key": "soja",
        "total_days": 7,
        "rain_mm": 60.0,
        "water_only": True,
        "solutes": [],
    }
    data.update(overrides)
    return SimulateRequest(**data)


def test_schema_defaults_water_only_off_and_solutes_optional() -> None:
    req = SimulateRequest(soil_key="latossolo_medio")
    assert req.water_only is False
    assert req.solutes == []


def test_simulate_water_only_empty_solutes_returns_water_balance() -> None:
    payload = {
        "soil_key": "latossolo_medio",
        "crop_key": "soja",
        "total_days": 7,
        "rain_mm": 60.0,
        "water_only": True,
        "solutes": [],
    }
    with TestClient(app) as client:
        response = client.post("/simulate", json=payload)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["water_only"] is True
    assert body["solutes"] == []
    # Galeria por-soluto vazia (Gradio: _solute_gallery_out = []).
    assert body["plots_png_base64"] == {}
    # Imagem principal é o perfil de água θ(z).
    assert base64.b64decode(body["plot_png_base64"]).startswith(PNG_MAGIC)
    wb = body["water_balance"]
    assert wb["applied_mm"] == 60.0
    assert wb["applied_label"] == "Lâmina Aplicada"
    assert wb["drainage_label"] == "Drenagem no Fundo"
    assert wb["drainage_mm"] == body["drainage_mm"]
    assert isinstance(wb["storage_mm"], float)
    assert isinstance(wb["runoff_mm"], float)
    assert isinstance(wb["evap_mm"], float)
    assert isinstance(wb["transp_mm"], float)


def test_simulate_nutrients_mode_without_solutes_returns_422() -> None:
    payload = {
        "soil_key": "latossolo_medio",
        "total_days": 7,
        "rain_mm": 60.0,
        "solutes": [],
    }
    with TestClient(app) as client:
        response = client.post("/simulate", json=payload)
    assert response.status_code == 422, response.text
    detail = response.json()["detail"]
    assert any(item.get("field") == "simulation_mode" for item in detail)


def test_water_balance_present_for_nutrient_runs() -> None:
    payload = {
        "soil_key": "latossolo_medio",
        "crop_key": "soja",
        "total_days": 7,
        "rain_mm": 60.0,
        "solutes": [{"symbol": "K", "name": "Potássio", "dose_kg_ha": 30.0, "kd": 0.45}],
    }
    with TestClient(app) as client:
        response = client.post("/simulate", json=payload)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["water_only"] is False
    assert len(body["solutes"]) == 1
    wb = body["water_balance"]
    assert wb["applied_mm"] == 60.0
    assert wb["applied_label"] == "Lâmina Aplicada"


def test_water_only_column_report_is_water_balance_only(tmp_path) -> None:
    _resp, _excel_path, report = execute_job(_water_only_request(), str(tmp_path))
    assert "Fluxo de Água" in report
    assert "Balanço Hídrico" in report
    assert "Kd" not in report
    assert "Dispersividade" not in report
    assert "lixivia" not in report.lower()
    # Passo a passo do core (water_only): Richards, sem Advecção-Dispersão.
    assert "Richards" in report
    assert "Advecção-Dispersão" not in report


def test_water_only_excel_summary_is_water_balance(tmp_path) -> None:
    _resp, excel_path, _report = execute_job(_water_only_request(), str(tmp_path))
    workbook = load_workbook(excel_path)
    sheet = workbook["Parametros_e_Balanco"]
    params = [str(row[0].value) for row in list(sheet.rows)[1:]]
    joined = " ".join(params)
    assert "Kd" not in joined
    assert "Dose" not in joined
    assert "water_only" not in joined
    assert "Modo de Simulação" in params
    assert any("Drenagem" in p for p in params)
    assert any("Armazenamento" in p for p in params)


def test_water_only_seasonal_report_and_excel(tmp_path) -> None:
    rain_series = [5.0 if day % 3 == 0 else 0.0 for day in range(365)]
    req = _water_only_request(rain_series=rain_series)
    _resp, excel_path, report = execute_job(req, str(tmp_path))
    assert "Fluxo de Água" in report
    assert "Kd" not in report
    assert "lixivia" not in report.lower()
    workbook = load_workbook(excel_path)
    sheet = workbook["Parametros_e_Balanco"]
    joined = " ".join(str(row[0].value) for row in list(sheet.rows)[1:])
    assert "water_only" not in joined
    assert "Kd [" not in joined


def test_water_only_seasonal_labels_annual(tmp_path) -> None:
    rain_series = [5.0 if day % 3 == 0 else 0.0 for day in range(365)]
    with TestClient(app) as client:
        response = client.post(
            "/simulate",
            json={
                "soil_key": "latossolo_medio",
                "crop_key": "soja",
                "rain_series": rain_series,
                "water_only": True,
                "solutes": [],
            },
        )
    assert response.status_code == 200, response.text
    wb = response.json()["water_balance"]
    assert wb["applied_mm"] == sum(rain_series)
    assert wb["applied_label"] == "Chuva Anual Total"
    assert wb["drainage_label"] == "Drenagem Anual"
