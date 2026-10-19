"""Testes de integração dos jobs sazonais e de coluna via API HTTP."""

import io
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

POLL_TIMEOUT_S = 240
POLL_INTERVAL_S = 0.5


def _base_payload(**overrides) -> dict:
    payload = {
        "soil_key": "latossolo_medio",
        "crop_key": "soja",
        "solutes": [
            {
                "symbol": "K",
                "name": "Potássio",
                "dose_kg_ha": 60.0,
                "kd": 1.0,
            }
        ],
    }
    payload.update(overrides)
    return payload


def _wait_for_done(client: TestClient, job_id: str) -> dict:
    deadline = time.monotonic() + POLL_TIMEOUT_S
    status = None
    while time.monotonic() < deadline:
        response = client.get(f"/jobs/{job_id}")
        assert response.status_code == 200, response.text
        status = response.json()
        if status["status"] in {"done", "failed"}:
            break
        time.sleep(POLL_INTERVAL_S)
    assert status is not None, "job não retornou status"
    assert status["status"] == "done", f"job terminou com status {status['status']}: {status.get('error')}"
    return status


def _assert_excel_and_report(client: TestClient, job_id: str) -> None:
    excel = client.get(f"/jobs/{job_id}/excel")
    assert excel.status_code == 200, excel.text
    workbook = load_workbook(io.BytesIO(excel.content))
    assert len(workbook.sheetnames) >= 1

    report = client.get(f"/jobs/{job_id}/report")
    assert report.status_code == 200, report.text
    assert report.json()["report_markdown"].strip()


def test_seasonal_job_365_days() -> None:
    rain_series = [5.0 if day % 3 == 0 else 0.0 for day in range(365)]
    payload = _base_payload(rain_series=rain_series)

    with TestClient(app) as client:
        created = client.post("/jobs", json=payload)
        assert created.status_code == 200, created.text
        job_id = created.json()["job_id"]

        status = _wait_for_done(client, job_id)
        assert status["status"] == "done"
        assert status["has_excel"] is True
        assert status["has_report"] is True

        _assert_excel_and_report(client, job_id)


def test_column_job_30_days_uniform_rain() -> None:
    payload = _base_payload(total_days=30, rain_mm=60.0)

    with TestClient(app) as client:
        created = client.post("/jobs", json=payload)
        assert created.status_code == 200, created.text
        job_id = created.json()["job_id"]

        status = _wait_for_done(client, job_id)
        assert status["status"] == "done"
        assert status["has_excel"] is True
        assert status["has_report"] is True

        _assert_excel_and_report(client, job_id)
